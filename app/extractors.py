"""Turn uploaded files into labelled text segments (Page N, Slide N, Sheet X ...)."""
import csv
import io
import shutil
import subprocess
import tempfile
from dataclasses import asdict, dataclass
from pathlib import Path

from . import ocr

MIN_PAGE_CHARS = 30  # below this a PDF page is treated as scanned and OCR'd

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp", ".gif"}
TEXT_EXTS = {".txt", ".md", ".markdown", ".rst", ".log", ".json"}
SUPPORTED_EXTS = (
    {".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".html", ".htm", ".doc", ".ppt", ".xls", ".odt", ".odp", ".rtf"}
    | IMAGE_EXTS
    | TEXT_EXTS
)


class ExtractionError(Exception):
    pass


@dataclass
class Segment:
    text: str
    source: str
    ocr: bool = False

    def to_dict(self) -> dict:
        return asdict(self)


def extract(path: Path) -> list[Segment]:
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTS:
        raise ExtractionError(f"Unsupported file type: {ext}")
    if ext == ".pdf":
        segs = _pdf(path)
    elif ext == ".docx":
        segs = _docx(path)
    elif ext == ".pptx":
        segs = _pptx(path)
    elif ext == ".xlsx":
        segs = _xlsx(path)
    elif ext == ".csv":
        segs = _csv(path)
    elif ext in (".html", ".htm"):
        segs = _html(path)
    elif ext in IMAGE_EXTS:
        segs = _image(path)
    elif ext in TEXT_EXTS:
        segs = _split_text(_read_text(path), "Section")
    else:  # legacy / office formats -> convert with LibreOffice if present
        segs = _legacy(path)
    segs = [s for s in segs if s.text.strip()]
    if not segs:
        hint = ""
        if not ocr.ocr_available():
            hint = " If this is a scanned document or image, install Tesseract OCR (see README)."
        raise ExtractionError(f"No readable text found in {path.name}.{hint}")
    return segs


# ---------- individual formats ----------

def _pdf(path: Path) -> list[Segment]:
    import pymupdf as fitz

    segs = []
    with fitz.open(path) as doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            used_ocr = False
            if len(text) < MIN_PAGE_CHARS:
                ocr_text = ocr.ocr_pdf_page(page)
                if len(ocr_text) > len(text):
                    text, used_ocr = ocr_text, True
            segs.append(Segment(text, f"Page {i}", used_ocr))
    return segs


def _docx(path: Path) -> list[Segment]:
    import docx

    d = docx.Document(str(path))
    segs: list[Segment] = []
    heading = "Introduction"
    buf: list[str] = []

    def flush():
        if buf:
            segs.append(Segment("\n".join(buf), heading))
            buf.clear()

    for para in d.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading") or style == "title":
            flush()
            heading = text[:80]
            buf.append(f"# {text}")
        else:
            buf.append(text)
    flush()

    for t_idx, table in enumerate(d.tables, start=1):
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        segs.append(Segment("\n".join(rows), f"Table {t_idx}"))

    # Long sections without headings -> split so citations stay useful
    out = []
    for s in segs:
        out.extend(_split_text(s.text, s.source) if len(s.text) > 6000 else [s])
    return out


def _pptx(path: Path) -> list[Segment]:
    from pptx import Presentation

    prs = Presentation(str(path))
    segs = []
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            parts.extend(_shape_text(shape))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Speaker notes: {notes}")
        segs.append(Segment("\n".join(p for p in parts if p), f"Slide {i}"))
    return segs


def _shape_text(shape) -> list[str]:
    parts = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        parts.append(shape.text_frame.text.strip())
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    if shape.shape_type == 6:  # group
        for s in shape.shapes:
            parts.extend(_shape_text(s))
    if shape.shape_type == 13 and ocr.ocr_available():  # picture -> OCR text inside images
        try:
            txt = ocr.ocr_image_bytes(shape.image.blob)
            if len(txt) > 20:
                parts.append(f"[Image text] {txt}")
        except Exception:
            pass
    return parts


def _xlsx(path: Path) -> list[Segment]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    segs = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                rows.append(" | ".join(cells))
        segs.extend(_split_text("\n".join(rows), f"Sheet {ws.title}"))
    wb.close()
    return segs


def _csv(path: Path) -> list[Segment]:
    text = _read_text(path)
    rows = [" | ".join(r) for r in csv.reader(io.StringIO(text))]
    return _split_text("\n".join(rows), "Rows")


def _html(path: Path) -> list[Segment]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_read_text(path), "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    return _split_text(soup.get_text("\n"), "Section")


def _image(path: Path) -> list[Segment]:
    if not ocr.ocr_available():
        raise ExtractionError(
            f"{path.name} is an image; reading it requires Tesseract OCR, which is not installed (see README)."
        )
    from PIL import Image

    with Image.open(path) as img:
        return [Segment(ocr.ocr_image(img), "Image", True)]


def _legacy(path: Path) -> list[Segment]:
    soffice = shutil.which("soffice") or shutil.which("soffice.exe")
    for candidate in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                      r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if not soffice and Path(candidate).exists():
            soffice = candidate
    target = {".doc": "docx", ".odt": "docx", ".rtf": "docx", ".ppt": "pptx", ".odp": "pptx", ".xls": "xlsx"}[
        path.suffix.lower()
    ]
    if not soffice:
        raise ExtractionError(
            f"{path.suffix} files can't be read directly. Open it and 'Save As' .{target}, "
            "or install LibreOffice for automatic conversion."
        )
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            [soffice, "--headless", "--convert-to", target, "--outdir", tmp, str(path)],
            check=True, capture_output=True, timeout=180,
        )
        converted = Path(tmp) / f"{path.stem}.{target}"
        if not converted.exists():
            raise ExtractionError(f"LibreOffice could not convert {path.name}.")
        return extract(converted)


# ---------- helpers ----------

def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _split_text(text: str, label: str, size: int = 4000) -> list[Segment]:
    """Split long plain text into numbered parts on paragraph boundaries."""
    text = text.strip()
    if not text:
        return []
    if len(text) <= size:
        return [Segment(text, label)]
    parts, buf = [], ""
    for para in text.split("\n"):
        if len(buf) + len(para) > size and buf:
            parts.append(buf)
            buf = ""
        buf += para + "\n"
    if buf.strip():
        parts.append(buf)
    return [Segment(p.strip(), f"{label} {i}") for i, p in enumerate(parts, start=1)]
