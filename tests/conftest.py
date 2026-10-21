import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

FIXTURES = Path(__file__).parent / "fixtures"

PHOTO_TEXT = "Photosynthesis converts light energy into chemical energy in chloroplasts."
MITO_TEXT = "The mitochondria is the powerhouse of the cell and produces ATP through respiration."


@pytest.fixture(scope="session")
def fixtures() -> dict[str, Path]:
    FIXTURES.mkdir(exist_ok=True)
    out = {}

    import pymupdf as fitz
    pdf = FIXTURES / "sample.pdf"
    doc = fitz.open()
    for text in (PHOTO_TEXT, MITO_TEXT):
        page = doc.new_page()
        page.insert_textbox(fitz.Rect(50, 50, 550, 300), text, fontsize=12)
    doc.save(pdf)
    doc.close()
    out["pdf"] = pdf

    import docx
    d = docx.Document()
    d.add_heading("Cell Biology", level=1)
    d.add_paragraph(MITO_TEXT)
    d.add_heading("Plants", level=1)
    d.add_paragraph(PHOTO_TEXT)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Organelle", "Function"
    t.cell(1, 0).text, t.cell(1, 1).text = "Ribosome", "Protein synthesis"
    out["docx"] = FIXTURES / "sample.docx"
    d.save(out["docx"])

    from pptx import Presentation
    prs = Presentation()
    for title, body in (("Photosynthesis", PHOTO_TEXT), ("Mitochondria", MITO_TEXT)):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = body
        s.notes_slide.notes_text_frame.text = f"Remember {title} for the exam"
    out["pptx"] = FIXTURES / "sample.pptx"
    prs.save(out["pptx"])

    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Glossary"
    wb.active.append(["Term", "Definition"])
    wb.active.append(["ATP", "Energy currency of the cell"])
    out["xlsx"] = FIXTURES / "sample.xlsx"
    wb.save(out["xlsx"])

    out["txt"] = FIXTURES / "sample.txt"
    out["txt"].write_text(PHOTO_TEXT + "\n\n" + MITO_TEXT, encoding="utf-8")

    out["html"] = FIXTURES / "sample.html"
    out["html"].write_text(f"<html><script>bad()</script><body><h1>Bio</h1><p>{MITO_TEXT}</p></body></html>",
                           encoding="utf-8")

    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (1400, 200), "white")
    try:
        font = ImageFont.truetype("arial.ttf", 40)
    except OSError:
        font = ImageFont.load_default()
    ImageDraw.Draw(img).text((20, 70), "Osmosis moves water across membranes", fill="black", font=font)
    out["png"] = FIXTURES / "sample.png"
    img.save(out["png"])

    # Scanned PDF: an image-only page
    scanned = fitz.open()
    page = scanned.new_page()
    page.insert_image(fitz.Rect(20, 20, 580, 120), filename=str(out["png"]))
    out["scanned_pdf"] = FIXTURES / "scanned.pdf"
    scanned.save(out["scanned_pdf"])
    scanned.close()
    return out
